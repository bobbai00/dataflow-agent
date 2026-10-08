# Licensed to the Apache Software Foundation (ASF) under one or more
# contributor license agreements. See the NOTICE file distributed with
# this work for additional information regarding copyright ownership.
# The ASF licenses this file to You under the Apache License, Version 2.0
# (the "License"); you may not use this file except in compliance with
# the License. You may obtain a copy of the License at
# http://www.apache.org/licenses/LICENSE-2.0
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Prepare twelve pinned DSBench analysis questions with separate gold metadata."""

import argparse
import hashlib
import io
import json
from pathlib import Path
import zipfile

import openpyxl


GITHUB_REVISION = "ba786096137a5108af11c016ad3f09cdb97beefd"
HF_REVISION = "1196d6553ec200e1262841efc89cda5954ffe89f"
ARCHIVE_URL = f"https://huggingface.co/datasets/liqiang888/DSBench/resolve/{HF_REVISION}/data_analysis/data.zip"
METADATA_URL = f"https://raw.githubusercontent.com/LiqiangJing/DSBench/{GITHUB_REVISION}/data_analysis/data.json"
ARCHIVE_SHA256 = "3a53d89bf0307978522d684fc2b49c8697a754444c0c2dd2e4771d0d8fdfb4e0"
METADATA_SHA256 = "ff822b69d46e08b384fa9b93ff39762c33b9887d315a5653bb536aeaba7572ea"

CHALLENGES = {
    "00000034": {
        "questions": [5, 10, 14],
        "workbook": "MO13-Data-Analysis-Data-Workbook.xlsx",
        "sheet": "Sheet1",
        "header": True,
        "skip_lines": 0,
        "required_columns": [
            "Date",
            "Invoice Number",
            "Item Code",
            "Sales Person",
            "Manager On Duty",
            "Sales Price Per Unit",
            "Cost per Unit",
            "Quantity",
            "Postal Code",
        ],
    },
    "00000032": {
        "questions": [1, 3, 4],
        "workbook": "MO14-Round-1-Dealing-With-Data-Workbook.xlsx",
        "sheet": "Usage",
        "header": False,
        "skip_lines": 0,
        "required_columns": ["usage_text"],
    },
    "00000006": {
        "questions": [17, 23, 25],
        "workbook": "MO16-Round-2-Sec-3-Tally-Up-1.xlsx",
        "sheet": "Data",
        "header": True,
        "skip_lines": 7,
        "required_columns": [
            "Voter ID",
            "Age",
            "District Code",
            "Red",
            "Orange",
            "Yellow",
            "Green",
            "Blue",
            "Purple",
            "Brown",
            "Black",
        ],
    },
    "00000010": {
        "questions": [5, 9, 15],
        "workbook": "MO2017-Finals-Sec-1-Ladder-Up.xlsx",
        "sheet": "Data",
        "header": True,
        "skip_lines": 0,
        "required_columns": [
            "Date",
            "Round",
            "Team_H",
            "Team_A",
            "Score_H",
            "Score_A",
            "Time_HG",
            "Time_AG",
            "Player_HG",
            "Player_AG",
        ],
    },
}


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def checked_source(path, expected, label):
    data = Path(path).read_bytes()
    if sha256(data) != expected:
        raise ValueError(f"{label} SHA-256 does not match the pinned source")
    return data


def inspect_workbook(data, config):
    """Parse actual cell values; Excel declared dimensions can include empty rows."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if config["sheet"] not in workbook.sheetnames:
            raise ValueError(f"missing selected sheet {config['sheet']!r}")
        sheet = workbook[config["sheet"]]
        rows = sheet.iter_rows(min_row=config["skip_lines"] + 1, values_only=True)
        if config["header"]:
            headers = list(next(rows, ()))
            missing = set(config["required_columns"]) - set(headers)
            if missing:
                raise ValueError(f"missing required columns in {config['workbook']}: {sorted(missing)}")
        else:
            headers = config["required_columns"]
        populated_rows = 0
        populated_columns = 0
        for row in rows:
            cells = [index for index, value in enumerate(row) if value is not None]
            if cells:
                populated_rows += 1
                populated_columns = max(populated_columns, max(cells) + 1)
        if not populated_rows:
            raise ValueError(f"no data rows in {config['workbook']}")
        if not config["header"] and populated_columns != 1:
            raise ValueError("expected a single unstructured Usage column")
        return {
            "sheet_names": workbook.sheetnames,
            "selected_sheet": config["sheet"],
            "header": config["header"],
            "skip_lines": config["skip_lines"],
            "required_columns": config["required_columns"],
            "populated_data_rows": populated_rows,
            "populated_columns": populated_columns,
        }
    finally:
        workbook.close()


def write_json(path, value, mode=0o644):
    with path.open("x", encoding="utf-8") as stream:
        json.dump(value, stream, indent=2, ensure_ascii=False)
        stream.write("\n")
    path.chmod(mode)


def prepare_dataset(archive, metadata, output):
    """Check pins/layouts before writing allowlisted inputs and private answer rows."""
    output = Path(output)
    for name in ("inputs", "evaluation", "manifest.json"):
        if (output / name).exists() or (output / name).is_symlink():
            raise FileExistsError(f"refusing to replace existing {output / name}")
    archive_data = checked_source(archive, ARCHIVE_SHA256, "archive")
    metadata_data = checked_source(metadata, METADATA_SHA256, "metadata")
    records = [json.loads(line) for line in metadata_data.decode("utf-8").splitlines() if line.strip()]
    by_id = {record["id"]: record for record in records}
    if len(by_id) != len(records):
        raise ValueError("duplicate challenge IDs in metadata")
    tasks, gold, content = [], [], {}
    with zipfile.ZipFile(io.BytesIO(archive_data)) as bundle:
        for challenge_id, config in CHALLENGES.items():
            source = by_id[challenge_id]
            prefix = f"data/{challenge_id}/"
            introduction = bundle.read(prefix + "introduction.txt")
            introduction.decode("utf-8")
            workbook = bundle.read(prefix + config["workbook"])
            layout = inspect_workbook(workbook, config)
            answers = dict(zip(source["questions"], source["answers"], strict=True))
            if len(answers) != len(source["questions"]):
                raise ValueError("duplicate question IDs in metadata")
            for number in config["questions"]:
                question_id = f"question{number}"
                question_file = question_id + ".txt"
                task_id = f"{challenge_id}-{question_id}"
                question = bundle.read(prefix + question_file)
                inputs = {
                    "introduction.txt": introduction,
                    question_file: question,
                    config["workbook"]: workbook,
                }
                content[task_id] = inputs
                tasks.append(
                    {
                        "task_id": task_id,
                        "challenge_id": challenge_id,
                        "challenge_name": source["name"],
                        "question_id": question_id,
                        "input_dir": f"inputs/{task_id}",
                        "introduction_file": "introduction.txt",
                        "question_file": question_file,
                        "workbook_file": config["workbook"],
                        "allowed_files": list(inputs),
                        "files": [
                            {"name": name, "source_member": prefix + name, "bytes": len(data), "sha256": sha256(data)}
                            for name, data in inputs.items()
                        ],
                        "source_url": source["url"],
                        "layout": layout,
                    }
                )
                gold.append(
                    {
                        "task_id": task_id,
                        "challenge_id": challenge_id,
                        "question_id": question_id,
                        "question": question.decode("utf-8"),
                        "answer": answers[question_id],
                    }
                )
    provenance = {
        "github_revision": GITHUB_REVISION,
        "huggingface_revision": HF_REVISION,
        "archive_url": ARCHIVE_URL,
        "archive_sha256": ARCHIVE_SHA256,
        "archive_bytes": len(archive_data),
        "metadata_url": METADATA_URL,
        "metadata_sha256": METADATA_SHA256,
        "metadata_bytes": len(metadata_data),
    }
    manifest = {
        "schema_version": 1,
        "benchmark": "DSBench data analysis",
        "selection": "12 preselected self-contained questions across four file-analysis challenges",
        "provenance": provenance,
        "task_count": len(tasks),
        "tasks": tasks,
    }
    output.mkdir(parents=True, exist_ok=True)
    (output / "inputs").mkdir()
    (output / "evaluation").mkdir(mode=0o700)
    for task in tasks:
        folder = output / task["input_dir"]
        folder.mkdir()
        for name, data in content[task["task_id"]].items():
            with (folder / name).open("xb") as stream:
                stream.write(data)
    write_json(output / "evaluation/gold.json", {"provenance": provenance, "tasks": gold}, mode=0o600)
    write_json(output / "manifest.json", manifest)
    verify_prepared_dataset(output)
    return manifest


def verify_prepared_dataset(output):
    """Reject extra input files, symlinks, changed content, and escaped task paths."""
    output = Path(output).resolve()
    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    inputs = output / "inputs"
    expected_dirs = {task["task_id"] for task in manifest["tasks"]}
    if {path.name for path in inputs.iterdir()} != expected_dirs:
        raise ValueError("unexpected input task directories")
    for task in manifest["tasks"]:
        folder = output / task["input_dir"]
        if folder.is_symlink() or folder.resolve().parent != inputs:
            raise ValueError("input task directory escaped its designated root")
        if {path.name for path in folder.iterdir()} != set(task["allowed_files"]):
            raise ValueError(f"unexpected input files for {task['task_id']}")
        for file in task["files"]:
            path = folder / file["name"]
            if path.is_symlink() or path.resolve().parent != folder:
                raise ValueError("input file escaped its designated task directory")
            if sha256(path.read_bytes()) != file["sha256"]:
                raise ValueError(f"input SHA-256 mismatch: {task['task_id']}/{file['name']}")
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive", type=Path, required=True)
    parser.add_argument("--metadata", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    manifest = prepare_dataset(args.archive, args.metadata, args.output)
    print(f"Prepared {manifest['task_count']} tasks: {args.output / 'manifest.json'}")


if __name__ == "__main__":
    main()
